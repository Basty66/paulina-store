import openpyxl
from openpyxl.styles import Font, Alignment
from datetime import datetime
import os
import subprocess
import sys

# Configuración
EXCEL_FILE = "paulinastore.xlsx"
OUTPUT_FILE = "paulinastore_con_fuente.xlsx"
FONT_NAME = "Libre Barcode EAN13 Text"
FONT_FILE = "LibreBarcodeEAN13Text-Regular.ttf"

def check_font_installed():
    """Verifica si la fuente está instalada en Windows"""
    fonts_dir = os.path.join(os.environ["WINDIR"], "Fonts")
    font_path = os.path.join(fonts_dir, FONT_FILE)
    return os.path.exists(font_path)

# Verificar si la fuente está instalada
if not check_font_installed():
    print("=" * 60)
    print("FUENTE NO INSTALADA")
    print("=" * 60)
    print()
    print("Para instalar la fuente:")
    print("1. Haga doble clic en el archivo", FONT_FILE)
    print("2. Haga clic en 'Instalar'")
    print("3. Ejecute este script nuevamente")
    print()
    print("Abriendo el archivo de fuente...")
    
    # Abrir el archivo de fuente
    try:
        os.startfile(FONT_FILE)
    except:
        subprocess.Popen(["start", FONT_FILE], shell=True)
    
    sys.exit(1)

print("=" * 60)
print("APLICANDO FUENTE DE CÓDIGO DE BARRAS AL EXCEL")
print("=" * 60)
print(f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
print(f"Fuente: {FONT_NAME}")
print("-" * 60)

# Leer Excel
wb = openpyxl.load_workbook(EXCEL_FILE)
ws = wb["Datos de origen"]

# Configurar ancho de columna para código de barras
ws.column_dimensions["B"].width = 20

# Crear estilo para la fuente de código de barras
barcode_font = Font(name=FONT_NAME, size=14)
header_font = Font(name="Calibri", size=11, bold=True)

# Aplicar formato a los encabezados
for col in range(1, 6):
    cell = ws.cell(row=1, column=col)
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Aplicar fuente de código de barras a la columna B
codigos_procesados = 0
for row_idx in range(2, ws.max_row + 1):
    codigo = ws.cell(row=row_idx, column=2).value
    
    if codigo is not None:
        # Convertir a string
        codigo_str = str(int(codigo)) if isinstance(codigo, (int, float)) else str(codigo)
        
        # Normalizar a 13 dígitos
        if len(codigo_str) == 12:
            codigo_str = "0" + codigo_str
        elif len(codigo_str) == 14:
            codigo_str = codigo_str[:13]
        
        # Aplicar fuente de código de barras
        cell = ws.cell(row=row_idx, column=2)
        cell.value = codigo_str
        cell.font = barcode_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
        codigos_procesados += 1

# Guardar Excel
wb.save(OUTPUT_FILE)

print(f"Códigos procesados: {codigos_procesados}")
print("-" * 60)
print(f"Archivo guardado como: {OUTPUT_FILE}")
print("=" * 60)
print("\nINSTRUCCIONES:")
print("1. Abra el archivo Excel generado")
print("2. Los códigos en la columna B se mostrarán como barras")
print("3. Si no se ven como barras, verifique que la fuente esté instalada")
print("=" * 60)
