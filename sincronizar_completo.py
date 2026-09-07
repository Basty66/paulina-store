import openpyxl
from openpyxl.drawing.image import Image as XLImage
import barcode
from barcode.writer import ImageWriter
from PIL import Image as PILImage
import io
import os
import json
from datetime import datetime

# Configuración
EXCEL_FILE = "paulinastore.xlsx"
OUTPUT_FILE = "paulinastore_con_codigos.xlsx"
REGISTRY_FILE = "codigos_generados.json"
TEMP_DIR = "temp_barcodes"

# Crear directorios
if not os.path.exists(TEMP_DIR):
    os.makedirs(TEMP_DIR)

# Cargar registro de códigos ya generados
if os.path.exists(REGISTRY_FILE):
    with open(REGISTRY_FILE, "r", encoding="utf-8") as f:
        registry = json.load(f)
else:
    registry = {"codigos": [], "ultima_actualizacion": None}

# Leer Excel
wb = openpyxl.load_workbook(EXCEL_FILE)
ws = wb["Datos de origen"]

# Agregar columna de código de barras
ws.cell(row=1, column=5, value="Código de Barras")

print("=" * 60)
print("SINCRONIZADOR DE CÓDIGOS DE BARRAS - VERSIÓN COMPLETA")
print("=" * 60)
print(f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
print(f"Archivo Excel: {EXCEL_FILE}")
print("-" * 60)

nuevos = 0
existentes = 0
errores = []

for row_idx in range(2, ws.max_row + 1):
    producto = ws.cell(row=row_idx, column=1).value
    codigo = ws.cell(row=row_idx, column=2).value
    precio_venta = ws.cell(row=row_idx, column=3).value
    
    if codigo is None:
        continue
    
    # Convertir a string y normalizar
    codigo_str = str(int(codigo)) if isinstance(codigo, (int, float)) else str(codigo)
    
    # Ajustar longitud según formato
    if len(codigo_str) == 12:
        codigo_str = "0" + codigo_str
    elif len(codigo_str) == 14:
        codigo_str = codigo_str[:13]
    elif len(codigo_str) != 13:
        errores.append(f"Fila {row_idx}: {producto} - código no válido ({len(codigo_str)} dígitos)")
        continue
    
    # Verificar si ya existe el código
    if codigo_str in registry["codigos"]:
        existentes += 1
        continue
    
    # Generar nuevo código de barras
    try:
        ean = barcode.get("ean13", code=codigo_str, writer=ImageWriter())
        
        # Guardar temporalmente
        temp_path = os.path.join(TEMP_DIR, f"temp_{row_idx}")
        saved_path = ean.save(temp_path)
        
        # Redimensionar imagen para Excel
        img = PILImage.open(saved_path)
        img = img.resize((200, 100), PILImage.LANCZOS)
        
        # Guardar en buffer
        buffer = io.BytesIO()
        img.save(buffer, format="PNG")
        buffer.seek(0)
        
        # Crear imagen de openpyxl
        xl_img = XLImage(buffer)
        xl_img.width = 150
        xl_img.height = 75
        
        # Insertar imagen en la celda E (columna 5)
        cell_ref = f"E{row_idx}"
        ws.add_image(xl_img, cell_ref)
        
        # También agregar el código como texto
        ws.cell(row=row_idx, column=5, value=codigo_str)
        
        # Agregar al registro
        registry["codigos"].append(codigo_str)
        
        nuevos += 1
        print(f"  NUEVO: {producto}")
        print(f"         Código: {codigo_str}")
        print(f"         Archivo: {saved_path}")
        
    except Exception as e:
        errores.append(f"Fila {row_idx}: {producto} - {str(e)}")

# Actualizar registro
registry["ultima_actualizacion"] = datetime.now().isoformat()
with open(REGISTRY_FILE, "w", encoding="utf-8") as f:
    json.dump(registry, f, ensure_ascii=False, indent=2)

# Guardar nuevo Excel
wb.save(OUTPUT_FILE)

# Limpiar archivos temporales
import shutil
if os.path.exists(TEMP_DIR):
    shutil.rmtree(TEMP_DIR)

# Resumen
print("-" * 60)
print("RESUMEN:")
print(f"  Códigos existentes (ya generados): {existentes}")
print(f"  Códigos nuevos generados e insertados: {nuevos}")
print(f"  Errores: {len(errores)}")

if errores:
    print("\nErrores:")
    for err in errores:
        print(f"  - {err}")

print("-" * 60)
print(f"Total de códigos en registro: {len(registry['codigos'])}")
print(f"Archivo guardado como: {OUTPUT_FILE}")
print("=" * 60)
